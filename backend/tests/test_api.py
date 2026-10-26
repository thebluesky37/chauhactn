import os
import sys

import boto3
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

moto = pytest.importorskip("moto")
from moto import mock_aws  # noqa: E402


def event(method, path, body=None, groups="[admin]", qs=None):
    import json

    return {
        "body": json.dumps(body) if body is not None else None,
        "queryStringParameters": qs,
        "requestContext": {
            "http": {"method": method, "path": path},
            "authorizer": {"jwt": {"claims": {"email": "a@b.c", "cognito:groups": groups}}},
        },
    }


@pytest.fixture
def aws(monkeypatch):
    monkeypatch.setenv("AWS_DEFAULT_REGION", "ap-southeast-1")
    monkeypatch.setenv("AWS_ACCESS_KEY_ID", "x")
    monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "x")
    monkeypatch.setenv("TABLE_NAME", "pm-main")
    monkeypatch.setenv("FILES_BUCKET", "files")
    with mock_aws():
        ddb = boto3.client("dynamodb")
        ddb.create_table(
            TableName="pm-main",
            BillingMode="PAY_PER_REQUEST",
            AttributeDefinitions=[
                {"AttributeName": n, "AttributeType": "S"} for n in ("PK", "SK", "GSI1PK", "GSI1SK")
            ],
            KeySchema=[{"AttributeName": "PK", "KeyType": "HASH"}, {"AttributeName": "SK", "KeyType": "RANGE"}],
            GlobalSecondaryIndexes=[
                {
                    "IndexName": "GSI1",
                    "KeySchema": [
                        {"AttributeName": "GSI1PK", "KeyType": "HASH"},
                        {"AttributeName": "GSI1SK", "KeyType": "RANGE"},
                    ],
                    "Projection": {"ProjectionType": "ALL"},
                }
            ],
        )
        boto3.client("s3", region_name="ap-southeast-1").create_bucket(
            Bucket="files", CreateBucketConfiguration={"LocationConstraint": "ap-southeast-1"}
        )
        from common import db

        db._table = None
        db._s3 = None
        yield


def call(mod, ev):
    import json

    r = mod.handler(ev, None)
    return r["statusCode"], (json.loads(r["body"]) if r["body"] else None)


def test_guest_visibility(aws):
    from admin_api import handler as admin
    from public_api import handler as public

    code, p = call(admin, event("POST", "/admin/projects", {"name": "A", "startDate": "2026-01-01", "plannedEndDate": "2026-12-31"}))
    assert code == 201 and p["guestVisible"] is False

    assert call(public, event("GET", "/public/projects"))[1]["projects"] == []
    assert call(public, event("GET", f"/public/projects/{p['id']}"))[0] == 404

    call(admin, event("PATCH", f"/admin/projects/{p['id']}", {"guestVisible": True}))
    assert len(call(public, event("GET", "/public/projects"))[1]["projects"]) == 1
    assert call(public, event("GET", f"/public/projects/{p['id']}"))[0] == 200

    call(admin, event("PATCH", f"/admin/projects/{p['id']}", {"guestVisible": False}))
    assert call(public, event("GET", "/public/projects"))[1]["projects"] == []
    assert call(public, event("GET", f"/public/projects/{p['id']}"))[0] == 404


def test_admin_requires_group(aws):
    from admin_api import handler as admin

    code, _ = call(admin, event("GET", "/admin/projects", groups="[]"))
    assert code == 403


def test_end_date_validation(aws):
    from admin_api import handler as admin

    code, _ = call(admin, event("POST", "/admin/projects", {"name": "A", "startDate": "2026-05-01", "plannedEndDate": "2026-01-01"}))
    assert code == 400


def test_export_workbook_single_sheet():
    import io

    from openpyxl import load_workbook

    from admin_api.export import build_workbook, export_filename

    view = {
        "project": {"id": "p1", "name": "Tòa nhà A", "status": "in_progress", "guestVisible": True,
                    "startDate": "2026-01-01", "plannedEndDate": "2026-12-31", "actualEndDate": None},
        "parties": [
            {"id": "a", "type": "subcontractor", "name": "Thầu X", "contactName": "Nam", "phone": "090", "email": "x@y.z"},
            {"id": "b", "type": "investor", "name": "CĐT 1"},
        ],
        "tasks": [{"id": "t", "partyId": "a", "name": "Đổ móng", "startDate": "2026-02-01", "endDate": "2026-03-01"}],
        "notes": [{"text": "Ghi chú\nnhiều dòng", "createdAt": "2026-02-02T03:00:00.000000Z"}],
        "uploads": [
            {"id": "u", "partyType": "subcontractor", "partyId": "a", "taskId": "t", "recipientIds": [],
             "sentDate": "2026-02-10", "periodFrom": "2026-02-01", "periodTo": "2026-02-09",
             "description": "Tiến độ", "adminNote": "", "hidden": True, "createdAt": "2026-02-10T01:00:00.000000Z",
             "files": [{"name": "a.xlsx", "key": "k", "size": 2097152, "version": 2},
                       {"name": "b.dwg", "key": "k2", "size": 100, "version": 1}]},
        ],
    }
    wb = load_workbook(io.BytesIO(build_workbook(view)))
    assert wb.sheetnames == ["Báo cáo dự án"]
    ws = wb.active
    assert set(ws.tables) == {"tblHangMuc", "tblLichSu4"}  # mỗi bên một bảng, chỉ bên có dữ liệu mới có bảng
    text = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert any("TÒA NHÀ A" in t for t in text)
    assert "Đã ẩn" in text  # bản ghi ẩn vẫn được xuất, có cột trạng thái
    assert sum(1 for t in text if t in ("a.xlsx", "b.dwg")) == 2  # mỗi file một dòng
    assert export_filename('A/B:"C"', __import__("datetime").datetime(2026, 3, 4)) == "ABC_2026-03-04.xlsx"


def test_export_empty_project():
    import io

    from openpyxl import load_workbook

    from admin_api.export import build_workbook

    view = {"project": {"id": "p", "name": "Rỗng", "status": "paused", "guestVisible": False,
                        "startDate": "2026-01-01", "plannedEndDate": "2026-02-01"},
            "parties": [], "tasks": [], "uploads": [], "notes": []}
    ws = load_workbook(io.BytesIO(build_workbook(view))).active
    assert not ws.tables


def test_export_file_links():
    import io

    from openpyxl import load_workbook

    from admin_api.export import build_workbook

    view = {"project": {"id": "p1", "name": "A", "status": "done", "guestVisible": False,
                        "startDate": "2026-01-01", "plannedEndDate": "2026-02-01"},
            "parties": [{"id": "a", "type": "investor", "name": "CĐT"}], "tasks": [], "notes": [],
            "uploads": [{"id": "u", "partyType": "investor", "partyId": "a", "sentDate": "2026-01-02",
                         "hidden": False, "createdAt": "2026-01-02T01:00:00.000000Z",
                         "files": [{"name": "bản vẽ.pdf", "key": "projects/p1/x/ban ve.pdf", "size": 10, "version": 1}]}]}
    ws = load_workbook(io.BytesIO(build_workbook(view, link_base="https://www.hachauctn.id.vn/"))).active
    links = [c.hyperlink.target for row in ws.iter_rows() for c in row if c.hyperlink]
    assert links == ["https://www.hachauctn.id.vn/download?p=p1&k=projects%2Fp1%2Fx%2Fban%20ve.pdf"]
    ws = load_workbook(io.BytesIO(build_workbook(view))).active
    assert not [c for row in ws.iter_rows() for c in row if c.hyperlink]


def test_delete_project_removes_everything(aws):
    from admin_api import handler as admin
    from common import db

    s3 = boto3.client("s3", region_name="ap-southeast-1")
    s3.put_bucket_versioning(Bucket="files", VersioningConfiguration={"Status": "Enabled"})

    def make_project(name):
        _, p = call(admin, event("POST", "/admin/projects", {"name": name, "startDate": "2026-01-01", "plannedEndDate": "2026-12-31"}))
        _, party = call(admin, event("POST", f"/admin/projects/{p['id']}/parties", {"type": "investor", "name": "CĐT"}))
        key = f"projects/{p['id']}/abc/f.pdf"
        s3.put_object(Bucket="files", Key=key, Body=b"1")
        s3.put_object(Bucket="files", Key=key, Body=b"22")  # 2 phiên bản
        _, up = call(admin, event("POST", f"/admin/projects/{p['id']}/uploads", {
            "partyType": "investor", "partyId": party["id"], "sentDate": "2026-02-01", "files": [{"name": "f.pdf", "key": key}]}))
        call(admin, event("PATCH", f"/admin/projects/{p['id']}/uploads/{up['id']}", {"hidden": True}))  # tạo REV
        call(admin, event("PUT", f"/admin/projects/{p['id']}/note", {"text": "x"}))
        return p["id"], up["id"]

    keep, _ = make_project("Giữ lại")
    gone, upload_id = make_project("Sẽ xóa")
    s3.put_object(Bucket="files", Key=f"exports/{gone}/r.xlsx", Body=b"x")
    s3.create_multipart_upload(Bucket="files", Key=f"projects/{gone}/dở/dang.pdf")

    # Không phải Admin thì không xóa được và không mất gì.
    status, _ = call(admin, event("DELETE", f"/admin/projects/{gone}", groups="[]"))
    assert status == 403
    assert db.get_project_item(gone) is not None

    status, body = call(admin, event("DELETE", f"/admin/projects/{gone}"))
    assert status == 200 and body["deleted"] == gone
    assert body["files"] == 3  # 2 phiên bản + báo cáo xuất

    assert db.get_project_item(gone) is None
    assert db.query_all(f"PROJECT#{gone}") == []
    assert db.query_all(f"UPLOAD#{upload_id}") == []  # lịch sử sửa
    leftovers = s3.list_object_versions(Bucket="files", Prefix=f"projects/{gone}/").get("Versions", [])
    assert leftovers == []
    assert s3.list_object_versions(Bucket="files", Prefix=f"exports/{gone}/").get("Versions", []) == []
    assert s3.list_multipart_uploads(Bucket="files").get("Uploads") is None

    # Dự án khác còn nguyên; nhật ký được giữ.
    assert db.get_project_item(keep) is not None
    assert len(db.query_all(f"PROJECT#{keep}")) > 1
    assert s3.list_object_versions(Bucket="files", Prefix=f"projects/{keep}/")["Versions"]
    assert any(a["action"] == "project.delete" for pk in db.table().scan()["Items"] for a in [pk] if "action" in a)

    assert call(admin, event("DELETE", f"/admin/projects/{gone}"))[0] == 404


def test_export_filtered_by_party_type():
    import io

    from openpyxl import load_workbook

    from admin_api.export import build_workbook

    view = {"project": {"id": "p", "name": "A", "status": "done", "guestVisible": False,
                        "startDate": "2026-01-01", "plannedEndDate": "2026-02-01"},
            "parties": [{"id": "a", "type": "investor", "name": "CĐT"},
                        {"id": "b", "type": "subcontractor", "name": "Thầu B"},
                        {"id": "c", "type": "architect", "name": "KTS C"}],
            "tasks": [{"id": "t", "partyId": "b", "name": "Móng", "startDate": "2026-01-02", "endDate": "2026-01-09"}],
            "notes": [],
            "uploads": [{"id": f"u{x}", "partyType": ty, "partyId": x, "sentDate": "2026-01-02", "hidden": False,
                         "createdAt": "2026-01-02T01:00:00.000000Z", "files": []}
                        for x, ty in (("a", "investor"), ("b", "subcontractor"), ("c", "architect"))]}
    ws = load_workbook(io.BytesIO(build_workbook(view, party_types=["investor", "subcontractor"]))).active
    text = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Thầu B" in text and "CĐT" in text and "KTS C" not in text
    assert "tblLichSu3" not in ws.tables  # kiến trúc sư (loại thứ 2) bị loại
    assert {"tblLichSu1", "tblLichSu4"} <= set(ws.tables) and "tblLichSu2" not in ws.tables


def test_party_scopes(aws):
    import io

    from openpyxl import load_workbook

    from admin_api import handler as admin
    from admin_api.export import build_workbook
    from common import db

    _, p = call(admin, event("POST", "/admin/projects", {"name": "P", "startDate": "2026-01-01", "plannedEndDate": "2026-12-31"}))
    base = f"/admin/projects/{p['id']}/parties"
    code, sub = call(admin, event("POST", base, {"type": "subcontractor", "name": "Thầu", "scopeIds": ["I.2", "I.1", "II.1"]}))
    assert code == 201 and sub["scopeIds"] == ["I.1", "I.2", "II.1"]  # theo thứ tự danh mục
    code, kts = call(admin, event("POST", base, {"type": "architect", "name": "KT", "scopeIds": ["I.1"]}))
    assert kts["scopeIds"] == []  # Kiến trúc không có scope
    code, _ = call(admin, event("POST", base, {"type": "investor", "name": "X", "scopeIds": ["Z.9"]}))
    assert code == 400
    code, upd = call(admin, event("PATCH", f"{base}/{sub['id']}", {
        "name": "Thầu mới", "scopeIds": ["VII.1", "I.1"], "scopeNotes": {"VII.1": " Ghi chú khí nén ", "I.1": "", "II.1": "bỏ"}}))
    assert code == 200 and upd["name"] == "Thầu mới" and upd["scopeIds"] == ["I.1", "VII.1"]
    assert upd["scopeNotes"] == {"VII.1": "Ghi chú khí nén"}  # bỏ ghi chú rỗng và ghi chú của scope không chọn
    code, upd = call(admin, event("PATCH", f"{base}/{sub['id']}", {"scopeIds": ["I.1"]}))
    assert upd["scopeNotes"] == {}  # bỏ chọn scope thì mất ghi chú
    code, upd = call(admin, event("PATCH", f"{base}/{sub['id']}", {"scopeIds": ["VII.1"], "scopeNotes": {"VII.1": "Ghi chú khí nén"}}))
    code, _ = call(admin, event("PATCH", f"{base}/nope", {"name": "x"}))
    assert code == 404

    ws = load_workbook(io.BytesIO(build_workbook(db.project_view(p["id"])))).active
    text = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "VII. HỆ THỐNG KHÍ NÉN" in text and "X" in text and "Ghi chú khí nén" in text
    assert "Ghi chú" in text and "KT\n(Kiến trúc)" not in text  # Kiến trúc không có cột scope
    assert "Thầu mới\n(Thầu phụ)" in text


def test_delete_party(aws):
    from admin_api import handler as admin

    _, p = call(admin, event("POST", "/admin/projects", {"name": "P", "startDate": "2026-01-01", "plannedEndDate": "2026-12-31"}))
    base = f"/admin/projects/{p['id']}/parties"
    _, free = call(admin, event("POST", base, {"type": "investor", "name": "Tự do"}))
    _, sub = call(admin, event("POST", base, {"type": "subcontractor", "name": "Thầu"}))
    call(admin, event("POST", f"/admin/projects/{p['id']}/tasks", {"partyId": sub["id"], "name": "Móng", "startDate": "2026-01-01", "endDate": "2026-02-01"}))

    code, _ = call(admin, event("DELETE", f"{base}/{sub['id']}"))
    assert code == 409  # còn hạng mục
    code, _ = call(admin, event("DELETE", f"{base}/{free['id']}"))
    assert code == 200
    code, _ = call(admin, event("DELETE", f"{base}/{free['id']}"))
    assert code == 404
    code, _ = call(admin, event("DELETE", f"{base}/{free['id']}", groups="[]"))
    assert code in (403, 404)


def test_scope_catalog_editable(aws):
    from admin_api import handler as admin

    code, cat = call(admin, event("GET", "/admin/scopes"))
    assert code == 200 and cat["groups"][0]["items"][0]["id"] == "I.1"  # chưa lưu thì dùng danh mục mặc định

    _, p = call(admin, event("POST", "/admin/projects", {"name": "P", "startDate": "2026-01-01", "plannedEndDate": "2026-12-31"}))
    base = f"/admin/projects/{p['id']}/parties"
    _, sub = call(admin, event("POST", base, {"type": "subcontractor", "name": "Thầu", "scopeIds": ["I.1", "II.1"], "scopeNotes": {"I.1": "a", "II.1": "b"}}))

    groups = cat["groups"]
    groups[0]["items"] = groups[0]["items"][1:]  # xóa I.1
    groups[0]["items"].append({"name": "  Scope mới  "})
    groups.append({"title": "Nhóm mới", "items": [{"name": "Mục 1"}]})
    code, saved = call(admin, event("PUT", "/admin/scopes", {"groups": groups}))
    assert code == 200 and saved["affectedParties"] == 1
    new_id = saved["groups"][0]["items"][-1]["id"]
    assert saved["groups"][0]["items"][-1]["name"] == "Scope mới" and new_id not in ("I.1", "")
    assert saved["groups"][-1]["id"].startswith("g")

    _, view = call(admin, event("GET", f"/admin/projects/{p['id']}"))
    party = view["parties"][0]
    assert party["scopeIds"] == ["II.1"] and party["scopeNotes"] == {"II.1": "b"}  # scope bị xóa được gỡ khỏi bên

    code, _ = call(admin, event("PATCH", f"{base}/{sub['id']}", {"scopeIds": [new_id]}))
    assert code == 200  # scope mới chọn được
    code, _ = call(admin, event("PATCH", f"{base}/{sub['id']}", {"scopeIds": ["I.1"]}))
    assert code == 400  # scope đã xóa không còn hợp lệ

    assert call(admin, event("PUT", "/admin/scopes", {"groups": [{"title": " ", "items": []}]}))[0] == 400
    assert call(admin, event("PUT", "/admin/scopes", {"groups": [{"title": "A", "items": [{"name": ""}]}]}))[0] == 400
    assert call(admin, event("PUT", "/admin/scopes", {"groups": "x"}))[0] == 400
    assert call(admin, event("GET", "/admin/scopes", groups="[]"))[0] == 403
